import shutil
from pathlib import Path

import pytest

pytest.importorskip("docx")
pytest.importorskip("docxcompose.composer")
pytest.importorskip("pypdf")

from docx import Document
from docx.enum.section import WD_SECTION
from pypdf import PdfReader, PdfWriter

from document_tools.features.merge import MergeOptions, MergeService
from document_tools.features.split import SplitMode, SplitOptions, SplitService
from document_tools.models import DocumentInput, OutputFormat


def _create_docx(path: Path, paragraphs: list[str]) -> None:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))


def _create_docx_with_sections(path: Path) -> None:
    document = Document()
    document.add_paragraph("Section 1 - Alpha")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("Section 2 - Beta")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("Section 3 - Gamma")
    document.save(str(path))


def _create_pdf(path: Path, page_count: int) -> None:
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=300, height=300)

    with path.open("wb") as file_obj:
        writer.write(file_obj)


def test_merge_docx_combines_paragraphs_in_order(tmp_path: Path) -> None:
    first_docx = tmp_path / "first.docx"
    second_docx = tmp_path / "second.docx"
    output_docx = tmp_path / "merged.docx"

    _create_docx(first_docx, ["Alpha", "Beta"])
    _create_docx(second_docx, ["Gamma", "Delta"])

    service = MergeService()
    result = service.merge(
        MergeOptions(
            inputs=[DocumentInput.from_path(first_docx), DocumentInput.from_path(second_docx)],
            output_path=output_docx,
            output_format=OutputFormat.DOCX,
        )
    )

    merged_document = Document(str(output_docx))
    merged_text = [paragraph.text for paragraph in merged_document.paragraphs if paragraph.text.strip()]

    assert result.success is True
    assert result.output_files == [output_docx]
    assert result.metadata["input_count"] == 2
    assert result.metadata["output_format"] == "docx"
    assert merged_text[:4] == ["Alpha", "Beta", "Gamma", "Delta"]


def test_split_docx_by_section_break_creates_one_file_per_section(tmp_path: Path) -> None:
    source_docx = tmp_path / "sections.docx"
    _create_docx_with_sections(source_docx)

    service = SplitService()
    result = service.split(
        SplitOptions(
            source=DocumentInput.from_path(source_docx),
            output_dir=tmp_path,
            mode=SplitMode.SECTION_BREAKS,
        )
    )

    assert result.success is True
    assert result.metadata["output_count"] == 3
    assert [path.name for path in result.output_files] == [
        "sections_section_0001.docx",
        "sections_section_0002.docx",
        "sections_section_0003.docx",
    ]

    first_text = [p.text for p in Document(str(result.output_files[0])).paragraphs if p.text.strip()]
    second_text = [p.text for p in Document(str(result.output_files[1])).paragraphs if p.text.strip()]
    third_text = [p.text for p in Document(str(result.output_files[2])).paragraphs if p.text.strip()]

    assert first_text == ["Section 1 - Alpha"]
    assert second_text == ["Section 2 - Beta"]
    assert third_text == ["Section 3 - Gamma"]


@pytest.mark.skipif(
    shutil.which("libreoffice") is None and shutil.which("soffice") is None,
    reason="LibreOffice is required for mixed DOCX+PDF to PDF conversion",
)
def test_merge_mixed_pdf_and_docx_to_pdf_creates_combined_output(tmp_path: Path) -> None:
    source_pdf = tmp_path / "first.pdf"
    source_docx = tmp_path / "second.docx"
    output_pdf = tmp_path / "mixed.pdf"

    _create_pdf(source_pdf, page_count=1)
    _create_docx(source_docx, ["Mixed document paragraph"])

    service = MergeService()
    result = service.merge(
        MergeOptions(
            inputs=[DocumentInput.from_path(source_pdf), DocumentInput.from_path(source_docx)],
            output_path=output_pdf,
            output_format=OutputFormat.PDF,
        )
    )

    merged_reader = PdfReader(str(output_pdf))

    assert result.success is True
    assert result.output_files == [output_pdf]
    assert result.metadata["input_count"] == 2
    assert result.metadata["converted_docx_count"] == 1
    assert len(merged_reader.pages) >= 2
